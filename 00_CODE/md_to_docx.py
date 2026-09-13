"""
Convert the PhD manuscript Markdown to a formatted Word (.docx) document.
Handles: headings (#..####), paragraphs, bold (**...**), italic (*...*),
bullet lists (- ...), horizontal rules (---), and GitHub-style tables.
Usage: python 00_CODE/md_to_docx.py
"""
from __future__ import annotations
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "04_DOCUMENTATION" / "MANUSCRIPT.md"
OUT = ROOT / "04_DOCUMENTATION" / "MANUSCRIPT.docx"

HEADING_SIZES = {1: 18, 2: 15, 3: 13, 4: 11.5}
ACCENT = RGBColor(0x0B, 0x5E, 0x6B)  # teal accent for headings


def add_runs_with_inline(paragraph, text):
    """Parse **bold** and *italic* inline markup into runs."""
    # Split on bold/italic tokens while keeping delimiters
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        else:
            paragraph.add_run(tok)


def is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_table_row(line: str):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_separator_row(line: str) -> bool:
    return bool(re.match(r"^\|?[\s:|-]+\|?$", line.strip())) and "-" in line


def main():
    md = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    n = len(md)
    while i < n:
        line = md[i]
        stripped = line.strip()

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            p = doc.add_paragraph()
            p_fmt = p.paragraph_format
            p_fmt.space_before = Pt(2)
            p_fmt.space_after = Pt(2)
            run = p.add_run("_" * 60)
            run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            i += 1
            continue

        # Heading
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            h = doc.add_heading(level=min(level, 4))
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = h.add_run("")
            add_runs_with_inline(h, text)
            for r in h.runs:
                r.font.size = Pt(HEADING_SIZES.get(level, 11))
                r.font.color.rgb = ACCENT
                r.bold = True
            i += 1
            continue

        # Table block
        if is_table_row(line) and i + 1 < n and is_separator_row(md[i + 1]):
            header = parse_table_row(line)
            i += 2  # skip header + separator
            rows = []
            while i < n and is_table_row(md[i]):
                rows.append(parse_table_row(md[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for j, cell_text in enumerate(header):
                hdr[j].paragraphs[0].text = ""
                add_runs_with_inline(hdr[j].paragraphs[0], cell_text)
                for r in hdr[j].paragraphs[0].runs:
                    r.bold = True
                    r.font.size = Pt(9.5)
            for row in rows:
                cells = table.add_row().cells
                for j in range(len(header)):
                    txt = row[j] if j < len(row) else ""
                    cells[j].paragraphs[0].text = ""
                    add_runs_with_inline(cells[j].paragraphs[0], txt)
                    for r in cells[j].paragraphs[0].runs:
                        r.font.size = Pt(9.5)
            doc.add_paragraph()
            continue

        # Bullet list
        if re.match(r"^[-*]\s+", stripped):
            text = re.sub(r"^[-*]\s+", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_inline(p, text)
            i += 1
            continue

        # Numbered list
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs_with_inline(p, m.group(1))
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        add_runs_with_inline(p, stripped)
        i += 1

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
